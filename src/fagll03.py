"""fagll03.py — Descarga las partidas abiertas de cuentas de mayor (FAGLL03).

Replica `resources/grabaciones/FAGLL03.vbs` y entrega el archivo tal como
lo exporta SAP, sin procesarlo. La grabación pega desde el portapapeles
dos listas: las cuentas de mayor y las sociedades.

Las cuentas salen del **maestro de cuentas** (Y_FI_GRL_004): las que tienen
`X` en la columna PI (gestión de partidas abiertas) para las sociedades
elegidas. Se usa el export del maestro más reciente de `salida/`, así que
primero hay que descargar el maestro. En el export de ejemplo, las 383
cuentas con partidas de FAGLL03 tienen todas PI = X en el maestro.

Diferencias con la grabación: `/n`, cuentas tomadas del maestro, sociedades
elegidas por el usuario, disposición por nombre (`DISPOSICION`) y archivo
en `salida/` como `FAGLL03_AAAAMMDD_HHMMSS.xlsx`.

Cumple el contrato de los módulos de flujo que usa la GUI:
`validar_configuracion()`, `get_sap_session()` y `descargar(...)`.

REQUISITOS
==========
Windows con SAP GUI abierto, sesión iniciada y scripting habilitado, y un
export del maestro (`salida/Y_FI_GRL_004_*.xlsx`).

USO
===
    python src/fagll03.py FECHA_CLAVE [SOCIEDAD ...]

    python src/fagll03.py 28.10.2026            # todas las sociedades
    python src/fagll03.py 28.10.2026 ISA XM     # solo ISA y XM
"""

from __future__ import annotations

import sys
from collections.abc import Callable, Iterable
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from comun import (
    SOCIEDADES,
    abrir_transaccion,
    buscar_ultimo_export,
    configurar_logging,
    crear_avisador,
    ejecutar_reporte,
    elegir_disposicion,
    elegir_fecha_en_calendario,
    exportar_hoja_calculo,
    get_sap_session,
    nombre_con_timestamp,
    pegar_seleccion_multiple,
    terminar_exportacion,
    validar_disposicion,
    validar_fecha_clave,
    validar_sociedades,
)
from paths import SALIDA_DIR
from y_fi_grl_004 import PREFIJO_SALIDA as PREFIJO_MAESTRO

# ---------------------------------------------------------------------------
# CONFIGURACIÓN
# ---------------------------------------------------------------------------

TRANSACCION = "FAGLL03"
PREFIJO_SALIDA = "FAGLL03"  # salida: FAGLL03_AAAAMMDD_HHMMSS.xlsx

# La misma disposición que FBL1N y FBL5N (confirmado por el equipo usuario
# el 2026-10-28). FAGLL03.vbs la elige por fila (473).
DISPOSICION: str | None = "/PARTABIERTA"

# IDs de FAGLL03.vbs.
ID_BTN_CUENTAS = "wnd[0]/usr/btn%_SD_SAKNR_%_APP_%-VALU_PUSH"
ID_BTN_SOCIEDADES = "wnd[0]/usr/btn%_SD_BUKRS_%_APP_%-VALU_PUSH"
ID_CAMPO_FECHA_CLAVE = "wnd[0]/usr/ctxtPA_STIDA"

# Columnas del export del maestro (disposición por defecto de
# Y_FI_GRL_004, ver docs/fichas-flujos.md).
COLUMNA_CUENTA = "Cta.mayor"
COLUMNA_PARTIDAS_ABIERTAS = "PI"
COLUMNA_SOCIEDAD = "Soc."
MARCA_PARTIDAS_ABIERTAS = "X"

# El export de ejemplo de todas las sociedades pesa ~34 MB (238 mil filas).
TIMEOUT_ARCHIVO_SEG = 600.0


# ---------------------------------------------------------------------------
# VALIDACIONES
# ---------------------------------------------------------------------------


def buscar_ultimo_maestro(carpeta: Path = SALIDA_DIR) -> Path | None:
    """El export del maestro más reciente en `carpeta`, o `None`."""
    return buscar_ultimo_export(PREFIJO_MAESTRO, carpeta)


def validar_maestro(carpeta: Path = SALIDA_DIR) -> Path:
    """Devuelve el maestro a usar o lanza `ValueError` si no hay ninguno."""
    maestro = buscar_ultimo_maestro(carpeta)
    if maestro is None:
        raise ValueError(
            "Primero descarga el Maestro de cuentas (Y_FI_GRL_004): las cuentas "
            f"de FAGLL03 salen de ese archivo y no hay ninguno en {carpeta}."
        )
    return maestro


def validar_configuracion() -> str:
    """Devuelve la disposición; lanza `ValueError` si falta la disposición o
    el export del maestro de cuentas."""
    disposicion = validar_disposicion(DISPOSICION, TRANSACCION)
    validar_maestro()
    return disposicion


# ---------------------------------------------------------------------------
# CUENTAS DEL MAESTRO (puro, openpyxl)
# ---------------------------------------------------------------------------


def leer_cuentas_partidas_abiertas(archivo_maestro: Path, sociedades: Iterable[str]) -> list[str]:
    """Cuentas de mayor con PI = X para alguna de `sociedades`, sin
    duplicados y en orden.

    La gestión de partidas abiertas se define por sociedad, por eso se
    filtra por la columna de sociedad del maestro.
    """
    elegidas = set(sociedades)
    wb = load_workbook(archivo_maestro, read_only=True, data_only=True)
    try:
        filas = wb.worksheets[0].iter_rows(values_only=True)
        encabezado = [str(c).strip() if c is not None else "" for c in next(filas, ())]
        requeridas = (COLUMNA_CUENTA, COLUMNA_PARTIDAS_ABIERTAS, COLUMNA_SOCIEDAD)
        faltantes = [c for c in requeridas if c not in encabezado]
        if faltantes:
            raise ValueError(
                f"El maestro {archivo_maestro.name} no tiene la(s) columna(s) "
                f"{', '.join(faltantes)}. Vuelve a descargarlo desde la app."
            )
        i_cuenta, i_pi, i_soc = (encabezado.index(c) for c in requeridas)
        cuentas = {
            str(fila[i_cuenta]).strip()
            for fila in filas
            if fila[i_cuenta]
            and str(fila[i_pi] or "").strip().upper() == MARCA_PARTIDAS_ABIERTAS
            and str(fila[i_soc] or "").strip().upper() in elegidas
        }
    finally:
        wb.close()
    if not cuentas:
        raise ValueError(
            f"El maestro {archivo_maestro.name} no tiene cuentas con PI = X "
            "para las sociedades elegidas."
        )
    return sorted(cuentas)


# ---------------------------------------------------------------------------
# PASOS: una función por paso; la sesión siempre es el 1.er argumento
# ---------------------------------------------------------------------------


def seleccionar_cuentas(session, cuentas: Iterable[str]) -> None:
    pegar_seleccion_multiple(session, ID_BTN_CUENTAS, cuentas)


def seleccionar_sociedades(session, sociedades: Iterable[str]) -> None:
    pegar_seleccion_multiple(session, ID_BTN_SOCIEDADES, sociedades)


def fijar_fecha_clave(session, fecha_clave: date) -> None:
    elegir_fecha_en_calendario(session, ID_CAMPO_FECHA_CLAVE, fecha_clave)


# ---------------------------------------------------------------------------
# ORQUESTADOR
# ---------------------------------------------------------------------------


def descargar(
    session,
    sociedades: Iterable[str],
    fecha_clave: date | str,
    carpeta_destino: Path = SALIDA_DIR,
    progreso: Callable[[str], None] | None = None,
    archivo_maestro: Path | None = None,
) -> tuple[str, str]:
    """Descarga FAGLL03 y devuelve `(carpeta, nombre_archivo)`.

    Las cuentas se leen del maestro (`archivo_maestro`, o el más reciente
    de `salida/`) ANTES de tocar SAP, igual que el resto de validaciones.
    Aborta en el primer error. `progreso(texto)` recibe el paso en curso.
    """
    socs = validar_sociedades(sociedades)
    fecha = validar_fecha_clave(fecha_clave)
    disposicion = validar_disposicion(DISPOSICION, TRANSACCION)
    maestro = Path(archivo_maestro) if archivo_maestro else validar_maestro()
    avisar = crear_avisador(progreso)

    avisar(f"Paso 1 de 8: leyendo las cuentas con PI = X de {maestro.name}…")
    cuentas = leer_cuentas_partidas_abiertas(maestro, socs)

    carpeta = Path(carpeta_destino)
    carpeta.mkdir(parents=True, exist_ok=True)
    nombre = nombre_con_timestamp(PREFIJO_SALIDA)

    avisar(
        f"FAGLL03: {len(cuentas)} cuenta(s), {len(socs)} sociedad(es), fecha clave {fecha:%d.%m.%Y}"
    )
    avisar("Paso 2 de 8: abriendo FAGLL03…")
    abrir_transaccion(session, TRANSACCION)
    avisar(f"Paso 3 de 8: cargando {len(cuentas)} cuentas y las sociedades…")
    seleccionar_cuentas(session, cuentas)
    seleccionar_sociedades(session, socs)
    avisar("Paso 4 de 8: fijando la fecha clave…")
    fijar_fecha_clave(session, fecha)
    avisar("Paso 5 de 8: ejecutando el reporte (puede tardar varios minutos)…")
    ejecutar_reporte(session)
    avisar(f"Paso 6 de 8: aplicando la disposición {disposicion}…")
    elegir_disposicion(session, disposicion)
    avisar("Paso 7 de 8: exportando a Excel (con todas las sociedades tarda)…")
    archivo = exportar_hoja_calculo(session, carpeta, nombre)
    avisar("Paso 8 de 8: esperando a que SAP termine de guardar…")
    terminar_exportacion(session, archivo, TIMEOUT_ARCHIVO_SEG)
    avisar(f"FAGLL03 listo: {archivo}")
    return str(carpeta), nombre


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    """0 = OK · 1 = error (validación o SAP) · 2 = uso incorrecto."""
    argv = argv if argv is not None else sys.argv[1:]
    if not argv:
        print("Uso: python src/fagll03.py FECHA_CLAVE [SOCIEDAD ...]", file=sys.stderr)
        return 2
    fecha_clave, sociedades = argv[0], argv[1:] or list(SOCIEDADES)
    configurar_logging()
    try:
        validar_sociedades(sociedades)
        validar_fecha_clave(fecha_clave)
        validar_configuracion()
        session = get_sap_session()
        carpeta, nombre = descargar(session, sociedades, fecha_clave)
    except (ValueError, RuntimeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"OK: {Path(carpeta) / nombre}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
